import io
from datetime import date, datetime, time

import pytest
from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from portal.imports import COLUMNS, FileRejected, build_template, parse_file, parse_time
from portal.models import Event, EventImport, Subscriber

IMPORT_URL = reverse("admin:portal_event_import")


def xlsx_file(rows, header=COLUMNS, name="events.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(name, output.getvalue())


def csv_file(text, name="events.csv"):
    return SimpleUploadedFile(name, text.encode("utf-8"))


def local(event_datetime):
    return timezone.localtime(event_datetime)


@pytest.fixture(autouse=True)
def media_in_tmp(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


# Parsing --------------------------------------------------------------------------------


@pytest.mark.django_db
def test_parse_xlsx_timed_all_day_and_bad_rows():
    upload = xlsx_file(
        [
            ["Science fair", date(2030, 10, 14), time(10, 0), None, time(14, 0), "Main hall", "Hi"],
            ["Diwali holidays", "20/10/2030", None, "24/10/2030", None, None, None],
            [None, None, None, None, None, None, None],  # blank rows are ignored
            ["Bad date", "31/02/2030", None, None, None, None, None],
            [None, "01/11/2030", None, None, None, None, None],
        ]
    )

    result = parse_file(upload)

    fair, holidays = result.events
    assert fair.title == "Science fair"
    assert local(fair.starts_at).time() == time(10, 0)
    assert local(fair.ends_at).time() == time(14, 0)
    assert not fair.all_day
    assert holidays.all_day
    assert local(holidays.starts_at).date() == date(2030, 10, 20)
    assert local(holidays.ends_at).date() == date(2030, 10, 24)
    assert result.errors == [
        "Row 5 (Bad date): '31/02/2030' isn't a date. Use a format like 14/10/2026.",
        "Row 6: the title is missing.",
    ]


@pytest.mark.django_db
def test_parse_csv_with_am_pm_times_and_loose_headings():
    upload = csv_file(
        "title*,START DATE,Start time,End time,Location\n"
        "Parents evening,2030-11-05,4:30 PM,7 PM,Classrooms\n"
    )

    [event] = parse_file(upload).events

    assert local(event.starts_at) == timezone.make_aware(datetime(2030, 11, 5, 16, 30))
    assert local(event.ends_at).time() == time(19, 0)
    assert event.location == "Classrooms"


@pytest.mark.parametrize(
    ("value", "expected"),
    [(0.5, time(12, 0)), ("9.30", time(9, 30)), ("2:15 p.m.", time(14, 15)), ("", None)],
)
def test_parse_time_formats(value, expected):
    assert parse_time(value) == expected


@pytest.mark.django_db
def test_end_before_start_is_reported():
    upload = xlsx_file([["Trip", "10/10/2030", "15:00", None, "09:00", None, None]])
    result = parse_file(upload)
    assert result.events == []
    assert result.errors == ["Row 2 (Trip): it ends before it starts."]


@pytest.mark.django_db
def test_duplicates_in_file_and_in_database_are_skipped():
    Event.objects.create(
        title="Sports day", starts_at=timezone.make_aware(datetime(2030, 9, 1, 9, 0))
    )
    upload = xlsx_file(
        [
            ["sports day", "01/09/2030", "13:00", None, None, None, None],
            ["Book fair", "02/09/2030", None, None, None, None, None],
            ["Book fair", "02/09/2030", None, None, None, None, None],
        ]
    )

    result = parse_file(upload)

    assert [e.title for e in result.events] == ["Book fair"]
    assert "already exists" in result.errors[0]
    assert "appears twice" in result.errors[1]


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("upload", "message"),
    [
        (SimpleUploadedFile("events.pdf", b"%PDF"), "Upload an Excel"),
        (xlsx_file([], header=["Name", "When"]), "Missing: Title, Start date"),
        (SimpleUploadedFile("events.xlsx", b"not really excel"), "couldn't be opened"),
    ],
)
def test_unusable_files_are_rejected(upload, message):
    with pytest.raises(FileRejected, match=message):
        parse_file(upload)


# All-day events -------------------------------------------------------------------------


@pytest.mark.django_db
def test_all_day_event_spans_whole_days_and_is_upcoming_on_the_day(client):
    today = timezone.localdate()
    event = Event.objects.create(
        title="Exam week",
        starts_at=timezone.make_aware(datetime.combine(today, time(15, 0))),
        ends_at=timezone.make_aware(datetime.combine(today, time(8, 0))),
        all_day=True,
    )

    assert local(event.starts_at).time() == time(0, 0)
    assert local(event.ends_at).time() == time(23, 59, 59)
    assert Event.objects.upcoming().filter(pk=event.pk).exists()
    assert b"Happening now" in client.get(reverse("portal:home")).content


# Admin upload and review ---------------------------------------------------------------


def upload_sample(admin_client):
    upload = xlsx_file(
        [
            ["Science fair", "14/10/2030", "10:00", None, "14:00", "Main hall", None],
            ["Diwali holidays", "20/10/2030", None, "24/10/2030", None, None, None],
            ["Broken", "not a date", None, None, None, None, None],
        ]
    )
    return admin_client.post(IMPORT_URL, {"file": upload, "language": "en"})


def review_post(batch, action, selected=(), **overrides):
    """Build the review form POST from the current drafts, with optional field overrides."""
    drafts = list(batch.events.filter(is_published=False).order_by("starts_at"))
    data = {
        "form-TOTAL_FORMS": str(len(drafts)),
        "form-INITIAL_FORMS": str(len(drafts)),
        "action": action,
        "selected": [str(e.pk) for e in selected],
    }
    for i, event in enumerate(drafts):
        start = local(event.starts_at)
        end = local(event.ends_at) if event.ends_at else None
        row = {
            "event_id": event.pk,
            "title": event.title,
            "start_date": start.date().isoformat(),
            "start_time": "" if event.all_day else start.strftime("%H:%M"),
            "end_date": end.date().isoformat() if end and end.date() != start.date() else "",
            "end_time": end.strftime("%H:%M") if end and not event.all_day else "",
            "location": event.location,
        }
        row.update(overrides.get(event.title, {}))
        data.update({f"form-{i}-{key}": value for key, value in row.items()})
    return data


def test_upload_creates_drafts_and_shows_review(admin_client):
    response = upload_sample(admin_client)

    batch = EventImport.objects.get()
    assert response.status_code == 302
    assert response.url == reverse("admin:portal_event_import_review", args=[batch.pk])
    assert batch.events.count() == 2
    assert not batch.events.filter(is_published=True).exists()
    assert len(batch.row_errors) == 1

    review = admin_client.get(response.url)
    assert b"Science fair" in review.content
    assert b"1 row couldn" in review.content


def test_drafts_are_not_public(admin_client, client):
    upload_sample(admin_client)
    assert b"Science fair" not in client.get(reverse("portal:event_list")).content


def test_file_with_no_valid_rows_creates_nothing(admin_client):
    upload = xlsx_file([["Broken", "nope", None, None, None, None, None]])
    response = admin_client.post(IMPORT_URL, {"file": upload, "language": "en"})
    assert response.status_code == 200
    assert b"No events could be imported" in response.content
    assert not EventImport.objects.exists()


def test_publish_ticked_events_with_email(admin_client, mailoutbox):
    Subscriber.objects.create(email="parent@example.com")
    upload_sample(admin_client)
    batch = EventImport.objects.get()
    fair = batch.events.get(title="Science fair")

    data = review_post(
        batch, "publish", selected=[fair], **{"Science fair": {"title": "Science fair 2030"}}
    )
    data["email_parents"] = "1"
    admin_client.post(reverse("admin:portal_event_import_review", args=[batch.pk]), data)

    fair.refresh_from_db()
    assert fair.is_published
    assert fair.title == "Science fair 2030"
    assert not batch.events.get(title="Diwali holidays").is_published
    assert len(mailoutbox) == 1


def test_publish_without_email_sends_nothing(admin_client, mailoutbox):
    Subscriber.objects.create(email="parent@example.com")
    upload_sample(admin_client)
    batch = EventImport.objects.get()

    admin_client.post(
        reverse("admin:portal_event_import_review", args=[batch.pk]),
        review_post(batch, "publish", selected=batch.events.all()),
    )

    assert batch.events.filter(is_published=True).count() == 2
    assert mailoutbox == []


def test_invalid_edit_blocks_saving(admin_client):
    upload_sample(admin_client)
    batch = EventImport.objects.get()
    fair = batch.events.get(title="Science fair")

    response = admin_client.post(
        reverse("admin:portal_event_import_review", args=[batch.pk]),
        review_post(batch, "publish", selected=[fair], **{"Science fair": {"end_time": "08:00"}}),
    )

    assert response.status_code == 200
    assert b"it ends before it starts" in response.content
    fair.refresh_from_db()
    assert not fair.is_published


def test_delete_ticked_drafts(admin_client):
    upload_sample(admin_client)
    batch = EventImport.objects.get()
    holidays = batch.events.get(title="Diwali holidays")

    admin_client.post(
        reverse("admin:portal_event_import_review", args=[batch.pk]),
        review_post(batch, "delete", selected=[holidays]),
    )

    assert list(batch.events.values_list("title", flat=True)) == ["Science fair"]


def test_cannot_publish_events_from_another_import(admin_client):
    upload_sample(admin_client)
    batch = EventImport.objects.get()
    other = Event.objects.create(title="Other draft", starts_at=timezone.now(), is_published=False)
    data = review_post(batch, "publish", selected=[other])

    admin_client.post(reverse("admin:portal_event_import_review", args=[batch.pk]), data)

    other.refresh_from_db()
    assert not other.is_published


def test_template_download(admin_client):
    response = admin_client.get(reverse("admin:portal_event_import_template"))
    workbook = load_workbook(io.BytesIO(response.content))
    assert [c.value for c in workbook["Events"][1]] == COLUMNS


@pytest.mark.django_db
def test_editor_can_import_but_public_user_cannot(client):
    call_command("setup_roles")
    editor = User.objects.create_user("teacher", password="pw-12345-long", is_staff=True)
    editor.groups.add(Group.objects.get(name="Editors"))

    assert client.get(IMPORT_URL).status_code == 302  # sent to the login page

    client.force_login(editor)
    assert client.get(IMPORT_URL).status_code == 200


def test_import_reads_the_type_in_three_languages(db):
    result = parse_file(
        csv_file(
            "Title,Start date,Type\n"
            "Cuti Deepavali,20/10/2026,  CUTI \n"
            "UPSA,12/10/2026,தேர்வு\n"
            "Sports day,15/10/2026,Sports\n"
            "Assembly,5/10/2026,\n"
        )
    )
    assert [e.kind for e in result.events] == ["holiday", "exam", "sports", "event"]
    assert result.errors == []


def test_unknown_type_becomes_event_with_a_note(db):
    result = parse_file(csv_file("Title,Start date,Type\nBook fair,20/10/2026,Pameran\n"))
    assert result.events[0].kind == "event"
    assert result.errors == [
        "Row 2 (Book fair): the type 'Pameran' wasn't recognised, so it was saved as Event."
    ]


def test_template_has_a_type_column():
    sheet = load_workbook(io.BytesIO(build_template())).active
    assert [cell.value for cell in sheet[1]][-1] == "Type"


def test_imported_type_is_saved_and_can_be_changed_on_review(admin_client):
    admin_client.post(
        reverse("admin:portal_event_import"),
        {
            "file": csv_file("Title,Start date,Type\nCuti Deepavali,20/10/2026,Cuti\n"),
            "language": "ms",
        },
    )
    event = Event.objects.get()
    assert event.kind == Event.HOLIDAY
    assert not event.is_published

    admin_client.post(
        reverse("admin:portal_event_import_review", args=[event.source_import_id]),
        {
            "form-TOTAL_FORMS": "1",
            "form-INITIAL_FORMS": "1",
            "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-event_id": str(event.pk),
            "form-0-title": "Cuti Deepavali",
            "form-0-start_date": "2026-10-20",
            "form-0-kind": Event.EXAM,
            "action": "save",
        },
    )
    event.refresh_from_db()
    assert event.kind == Event.EXAM
