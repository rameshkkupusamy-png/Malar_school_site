"""Adding people to the staff list (who may sign in with Google) from a spreadsheet.

Used by the "Import from spreadsheet" page on Staff members and by `manage.py import_staff`.
Each row is one person: the first email address in the row is used, and the first other
filled-in cell as their name. People already on the list are left as they are.
"""

import csv
import io
from dataclasses import dataclass, field
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import StaffMember

MAX_BYTES = 2 * 1024 * 1024


class StaffFileRejected(Exception):
    """The file can't be read at all; the message is shown to staff as-is."""


@dataclass
class StaffImportResult:
    added: list[str] = field(default_factory=list)
    already_listed: list[str] = field(default_factory=list)
    rows_without_email: list[int] = field(default_factory=list)


def is_email(value: str) -> bool:
    try:
        validate_email(value)
    except ValidationError:
        return False
    return True


def read_rows(fileobj, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            sheet = load_workbook(fileobj, read_only=True, data_only=True).active
        except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
            raise StaffFileRejected(
                "This file couldn't be opened as an Excel spreadsheet. Save it as .xlsx or CSV "
                "and try again."
            ) from error
        return [
            ["" if c is None else str(c).strip() for c in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    if name.endswith(".csv"):
        text = io.TextIOWrapper(fileobj, encoding="utf-8-sig", errors="replace", newline="")
        return [[c.strip() for c in row] for row in csv.reader(text)]
    raise StaffFileRejected("Upload an Excel (.xlsx) or CSV file.")


def add_staff_from_rows(rows: list[list[str]]) -> StaffImportResult:
    result = StaffImportResult()
    for number, row in enumerate(rows, start=1):
        cells = [c for c in row if c]
        if not cells:
            continue
        email = next((c.lower() for c in cells if is_email(c)), None)
        if email is None:
            result.rows_without_email.append(number)
            continue
        name = next((c for c in cells if c.lower() != email and "@" not in c), "")
        _, created = StaffMember.objects.get_or_create(email=email, defaults={"name": name})
        (result.added if created else result.already_listed).append(email)
    return result
